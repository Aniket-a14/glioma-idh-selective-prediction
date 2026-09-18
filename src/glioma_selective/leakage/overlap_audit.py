"""Build the Phase 03 patient-level overlap matrix.

The aggregate report is intentionally identifier-free: it records counts,
source coverage, and evidence for each matrix cell, but never a patient roster.
The optional detail report is a derived operational artifact and is written
under :func:`glioma_selective.utils.paths.manifests_dir`; it is not a Git
tracked manifest.

The UCSF-PDGM/BraTS-2021 comparison is special.  The two cohorts use different
identifier namespaces, so the official crosswalk and the UCSF metadata's
``BraTS21 ID`` column are compared through a source-to-BraTS mapping rather
than by intersecting the two canonical strings directly.  Both source views
are retained in the aggregate report when they disagree.
"""

from __future__ import annotations

import csv
import json
import re
from collections.abc import Iterable, Mapping
from pathlib import Path
from typing import Final, cast

from openpyxl import load_workbook  # type: ignore[import-untyped]

from glioma_selective.data.adapters import ucsf_pdgm, utsw
from glioma_selective.leakage.id_normalization import (
    Cohort,
    IdentifierNormalizationError,
    _coerce_cohort,
    normalize_identifier,
    normalize_ucsf_pdgm_id,
)
from glioma_selective.utils import paths

__all__ = [
    "BRATS_TEST_SET_CAVEAT",
    "EGD_NOT_AVAILABLE_REASON",
    "MATRIX_COLUMNS",
    "MATRIX_ROWS",
    "OverlapAuditError",
    "build_overlap_matrix",
    "build_real_overlap_audit",
    "write_overlap_detail_report",
    "write_overlap_report",
]


MATRIX_ROWS: Final[tuple[Cohort, ...]] = (Cohort.UTSW, Cohort.UCSF_PDGM, Cohort.EGD)
MATRIX_COLUMNS: Final[tuple[Cohort, ...]] = (
    Cohort.BRATS_2021,
    Cohort.TCGA_GBM,
    Cohort.TCGA_LGG,
    Cohort.UCSF_PDGM,
    Cohort.EGD,
)
BRATS_TEST_SET_CAVEAT: Final[str] = (
    "BraTS-2021 coverage is limited to the official training+validation "
    "crosswalk available here; the challenge test set is sequestered on "
    "Synapse (syn25829067) and is not obtainable, so this cell is not an "
    "exhaustive BraTS-2021 overlap result."
)
EGD_NOT_AVAILABLE_REASON: Final[str] = (
    "EGD has not been downloaded because its data-use agreement is pending "
    "(D-018); every EGD-involving overlap is unknowable until the governed "
    "inventory arrives."
)
_BRATS_DATA_COLLECTION = "Data Collection (as on TCIA+additional)"
_BRATS_PATIENT_ID = "PatientID on TCIA Radiology Portal"
_BRATS_ID = "BraTS2021 ID"
_UCSF_BRATS_COLLECTIONS: Final[frozenset[str]] = frozenset(
    {"UCSF-PDGM", "UCSF-PDGM_Additional"}
)
_TCGA_BARCODE_RE = re.compile(r"TCGA-[A-Z0-9]{2}-[A-Z0-9]{4}")


class OverlapAuditError(RuntimeError):
    """Raised when a raw overlap-audit source cannot be validated."""


def _cohort(value: Cohort | str) -> Cohort:
    """Normalize a public cohort argument through the existing dispatcher."""

    return _coerce_cohort(value)


def _normalize_set(cohort: Cohort, raw_ids: Iterable[str]) -> frozenset[str]:
    """Normalize an iterable and reject duplicate canonical IDs."""

    normalized: list[str] = []
    for raw_id in raw_ids:
        try:
            normalized.append(normalize_identifier(raw_id, cohort))
        except IdentifierNormalizationError as error:
            raise OverlapAuditError(
                f"cannot normalize {cohort.value} identifier for overlap audit"
            ) from error
    return frozenset(normalized)


def _normalize_ucsf_inventory_ids(
    raw_ids: Iterable[str],
) -> tuple[frozenset[str], dict[str, object]]:
    """Normalize UCSF directory IDs, including the observed follow-up wrapper.

    ``normalize_ucsf_pdgm_id`` remains the only function that canonicalizes the
    patient identifier.  The six observed ``..._FU..._nifti`` directory names
    are an acquisition wrapper around the base patient ID; they are collapsed
    to that base before calling the existing normalizer.
    """

    normalized: list[str] = []
    raw_count = 0
    followup_count = 0
    for raw_id in raw_ids:
        raw_count += 1
        try:
            normalized.append(normalize_ucsf_pdgm_id(raw_id))
            continue
        except IdentifierNormalizationError:
            followup = re.fullmatch(r"(UCSF-PDGM-\d{3,4})_FU[A-Za-z0-9]+_nifti", raw_id)
            if followup is None:
                raise OverlapAuditError(
                    "UCSF directory identifier is outside the T03.1 format "
                    "and the observed follow-up wrapper"
                ) from None
            followup_count += 1
            normalized.append(normalize_ucsf_pdgm_id(followup.group(1)))
    return frozenset(normalized), {
        "raw_directory_count": raw_count,
        "unique_normalized_patient_count": len(set(normalized)),
        "followup_directory_count": followup_count,
        "format_finding": (
            "Six UCSF directories use an observed _FU..._nifti follow-up "
            "wrapper. The wrapper was removed only to reuse the existing "
            "T03.1 base-ID normalizer; no new patient-ID padding rule was added."
        ),
    }


def _normalize_brats_source_id(raw_id: str) -> tuple[str, bool]:
    """Normalize a crosswalk BraTS ID and flag the observed bare-ID wrapper."""

    try:
        return normalize_identifier(raw_id, Cohort.BRATS_2021), False
    except IdentifierNormalizationError:
        if re.fullmatch(r"\d{5}", raw_id) is None:
            raise OverlapAuditError(
                "BraTS crosswalk identifier is outside the T03.1 format and "
                "the observed bare five-digit wrapper"
            ) from None
        return normalize_identifier(f"BraTS2021_{raw_id}", Cohort.BRATS_2021), True


def _pair_key(left: Cohort, right: Cohort) -> str:
    """Return a stable JSON key for a directed matrix cell."""

    return f"{left.value}|{right.value}"


def _pair_details_for_intersection(
    left: frozenset[str], right: frozenset[str]
) -> tuple[tuple[str, str], ...]:
    """Return equal normalized pairs without exposing them in the summary."""

    return tuple((identifier, identifier) for identifier in sorted(left & right))


def _cell(
    *,
    left: Cohort,
    right: Cohort,
    status: str,
    overlap_count: int | None,
    comparison_basis: str | None = None,
    reason: str | None = None,
    brats_caveat: bool = False,
) -> dict[str, object]:
    """Construct one complete, JSON-serializable matrix cell."""

    result: dict[str, object] = {
        "status": status,
        "overlap_count": overlap_count,
    }
    if comparison_basis is not None:
        result["comparison_basis"] = comparison_basis
    if reason is not None:
        result["reason"] = reason
    if brats_caveat:
        result["brats_coverage_caveat"] = BRATS_TEST_SET_CAVEAT
    if left == Cohort.UCSF_PDGM and right == Cohort.BRATS_2021:
        result["must_resolve_before_external_use"] = True
    return result


def build_overlap_matrix(
    identifiers: Mapping[Cohort | str, Iterable[str]],
    *,
    pair_details: Mapping[tuple[Cohort | str, Cohort | str], Iterable[tuple[str, str]]]
    | None = None,
    source_notes: Mapping[str, object] | None = None,
) -> dict[str, object]:
    """Build every cell in the Phase 03 3-by-5 overlap matrix.

    Args:
        identifiers: Raw or already canonical IDs for available cohorts.
            Values are normalized with the existing T03.1 functions.
        pair_details: Optional authoritative crosswalk pairs for a cell whose
            cohorts use different identifier namespaces.  The UCSF-PDGM and
            BraTS-2021 cell uses this mechanism.
        source_notes: Aggregate-only source evidence to include in the report.

    Cells involving EGD are ``not_yet_available`` when EGD is absent.  The two
    same-cohort cells represented by the requested matrix are
    ``not_applicable``.  All other cells are checked and carry an integer count,
    including zero.
    """

    normalized: dict[Cohort, frozenset[str]] = {}
    for raw_cohort, raw_ids in identifiers.items():
        selected = _cohort(raw_cohort)
        if selected in normalized:
            raise OverlapAuditError(f"duplicate cohort input: {selected.value}")
        normalized[selected] = _normalize_set(selected, raw_ids)

    normalized_pairs: dict[tuple[Cohort, Cohort], tuple[tuple[str, str], ...]] = {}
    if pair_details is not None:
        for (raw_left, raw_right), pairs in pair_details.items():
            key = (_cohort(raw_left), _cohort(raw_right))
            if key in normalized_pairs:
                raise OverlapAuditError(f"duplicate pair input: {_pair_key(*key)}")
            normalized_pairs[key] = tuple(sorted(pairs))

    matrix: dict[str, dict[str, dict[str, object]]] = {}
    detail_pairs: dict[str, list[tuple[str, str]]] = {}
    for left in MATRIX_ROWS:
        row: dict[str, dict[str, object]] = {}
        for right in MATRIX_COLUMNS:
            if left == right:
                row[right.value] = _cell(
                    left=left,
                    right=right,
                    status="not_applicable",
                    overlap_count=None,
                    reason="The row and column refer to the same cohort.",
                    brats_caveat=right == Cohort.BRATS_2021,
                )
                continue
            if left == Cohort.EGD or right == Cohort.EGD:
                row[right.value] = _cell(
                    left=left,
                    right=right,
                    status="not_yet_available",
                    overlap_count=None,
                    reason=EGD_NOT_AVAILABLE_REASON,
                    brats_caveat=right == Cohort.BRATS_2021,
                )
                continue

            key = (left, right)
            if key in normalized_pairs:
                pairs = normalized_pairs[key]
                basis = "authoritative_crosswalk_pair_intersection"
            elif (right, left) in normalized_pairs:
                pairs = tuple(
                    (right_id, left_id)
                    for left_id, right_id in normalized_pairs[(right, left)]
                )
                pairs = tuple(sorted(pairs))
                basis = "authoritative_crosswalk_pair_intersection"
            elif left not in normalized or right not in normalized:
                row[right.value] = _cell(
                    left=left,
                    right=right,
                    status="not_yet_available",
                    overlap_count=None,
                    reason=(
                        f"No identifier inventory is available for {left.value} "
                        f"or {right.value}."
                    ),
                    brats_caveat=right == Cohort.BRATS_2021,
                )
                continue
            else:
                pairs = _pair_details_for_intersection(
                    normalized[left], normalized[right]
                )
                basis = "normalized_identifier_intersection"

            row[right.value] = _cell(
                left=left,
                right=right,
                status="checked",
                overlap_count=len(pairs),
                comparison_basis=basis,
                brats_caveat=right == Cohort.BRATS_2021,
            )
            detail_pairs[_pair_key(left, right)] = list(pairs)
        matrix[left.value] = row

    report: dict[str, object] = {
        "report_version": 1,
        "matrix_rows": [cohort.value for cohort in MATRIX_ROWS],
        "matrix_columns": [cohort.value for cohort in MATRIX_COLUMNS],
        "matrix": matrix,
        "detail_pair_keys": sorted(detail_pairs),
    }
    if source_notes is not None:
        report["source_notes"] = dict(source_notes)
    return report


def _json_object(value: object, context: str) -> dict[str, object]:
    """Validate a decoded JSON object without introducing ``Any``."""

    if not isinstance(value, dict):
        raise OverlapAuditError(f"{context} must be a JSON object")
    return cast(dict[str, object], value)


def _load_tcga_ids(
    path: Path, cohort: Cohort
) -> tuple[frozenset[str], dict[str, object]]:
    """Load and normalize the raw GDC case response for one TCGA cohort."""

    try:
        payload = _json_object(json.loads(path.read_text(encoding="utf-8")), path.name)
    except (OSError, json.JSONDecodeError) as error:
        raise OverlapAuditError(f"cannot read TCGA source {path}") from error
    data = _json_object(payload.get("data"), f"{path.name}.data")
    hits = data.get("hits")
    if not isinstance(hits, list):
        raise OverlapAuditError(f"{path.name}.data.hits must be an array")
    raw_ids: list[str] = []
    format_matches = 0
    for index, value in enumerate(hits):
        hit = _json_object(value, f"{path.name}.data.hits[{index}]")
        submitter_id = hit.get("submitter_id")
        if not isinstance(submitter_id, str) or not submitter_id.strip():
            raise OverlapAuditError(f"{path.name} has an invalid submitter_id")
        raw_id = submitter_id.strip()
        raw_ids.append(raw_id)
        format_matches += int(_TCGA_BARCODE_RE.fullmatch(raw_id) is not None)
    normalized = _normalize_set(cohort, raw_ids)
    return normalized, {
        "cohort": cohort.value,
        "raw_case_count": len(raw_ids),
        "unique_normalized_count": len(normalized),
        "observed_format": (
            "TCGA-XX-XXXX barcode"
            if format_matches == len(raw_ids)
            else "mixed_or_unexpected"
        ),
        "normalizer_format_finding": (
            "The real GDC IDs are TCGA-XX-XXXX barcodes. T03.1's conservative "
            "generic normalizer prepends the cohort label; this is a format "
            "surprise to review, not a new interpretation made by T03.4."
        ),
    }


def _worksheet_rows(path: Path) -> tuple[tuple[object, ...], ...]:
    """Read the official crosswalk values without writing or modifying it."""

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (OSError, ValueError) as error:
        raise OverlapAuditError(f"cannot read BraTS crosswalk {path}") from error
    if "TCIA" not in workbook.sheetnames:
        raise OverlapAuditError("BraTS crosswalk has no TCIA sheet")
    worksheet = workbook["TCIA"]
    rows = tuple(
        tuple(cast(Iterable[object], row))
        for row in worksheet.iter_rows(values_only=True)
    )
    if not rows:
        raise OverlapAuditError("BraTS crosswalk is empty")
    return rows


def _load_brats_crosswalk(
    path: Path,
) -> tuple[
    frozenset[str],
    dict[tuple[str, str], tuple[str, str]],
    dict[str, object],
]:
    """Load BraTS IDs and UCSF source-to-BraTS mapping evidence."""

    rows = _worksheet_rows(path)
    header = rows[0]
    indexes = {
        str(value): index for index, value in enumerate(header) if value is not None
    }
    required = (_BRATS_DATA_COLLECTION, _BRATS_PATIENT_ID, _BRATS_ID)
    missing = [name for name in required if name not in indexes]
    if missing:
        raise OverlapAuditError(
            f"BraTS crosswalk is missing required columns: {', '.join(missing)}"
        )

    brats_ids: list[str] = []
    bare_brats_id_count = 0
    ucsf_primary: list[str] = []
    ucsf_additional: list[str] = []
    ucsf_primary_pairs: dict[tuple[str, str], tuple[str, str]] = {}
    for row in rows[1:]:
        collection = str(row[indexes[_BRATS_DATA_COLLECTION]] or "").strip()
        raw_brats_id = row[indexes[_BRATS_ID]]
        if raw_brats_id is not None and str(raw_brats_id).strip():
            brats_id, was_bare = _normalize_brats_source_id(str(raw_brats_id).strip())
            brats_ids.append(brats_id)
            bare_brats_id_count += int(was_bare)
        if collection not in _UCSF_BRATS_COLLECTIONS or raw_brats_id is None:
            continue
        brats_id, _ = _normalize_brats_source_id(str(raw_brats_id).strip())
        if collection == "UCSF-PDGM_Additional":
            ucsf_additional.append(brats_id)
            continue
        ucsf_primary.append(brats_id)
        raw_patient = row[indexes[_BRATS_PATIENT_ID]]
        if isinstance(raw_patient, bool) or not isinstance(raw_patient, int | str):
            raise OverlapAuditError(
                "UCSF-PDGM crosswalk patient IDs must be numeric or textual"
            )
        patient_text = str(raw_patient).strip()
        if not patient_text.isdigit() or len(patient_text) > 4:
            raise OverlapAuditError(
                "UCSF-PDGM crosswalk patient IDs must be numeric and at most "
                "four digits"
            )
        source_id = normalize_ucsf_pdgm_id(f"UCSF-PDGM-{patient_text.zfill(3)}")
        ucsf_primary_pairs[(source_id, brats_id)] = (source_id, brats_id)

    normalized_brats = frozenset(brats_ids)
    primary_set = frozenset(ucsf_primary)
    additional_set = frozenset(ucsf_additional)
    return (
        normalized_brats,
        ucsf_primary_pairs,
        {
            "crosswalk_row_count": len(rows) - 1,
            "crosswalk_unique_brats_id_count": len(normalized_brats),
            "bare_five_digit_brats_id_count": bare_brats_id_count,
            "brats_format_finding": (
                "Eight crosswalk rows contain a bare five-digit BraTS ID. The "
                "BraTS2021_ wrapper was added only to reuse the existing T03.1 "
                "normalizer; no new width or padding rule was added."
            ),
            "ucsf_pdgm_row_count": len(primary_set),
            "ucsf_pdgm_additional_row_count": len(additional_set),
            "ucsf_pdgm_all_row_count": len(primary_set | additional_set),
            "ucsf_pdgm_primary_brats_ids": primary_set,
            "ucsf_pdgm_additional_brats_ids": additional_set,
            "ucsf_pdgm_all_brats_ids": primary_set | additional_set,
        },
    )


def _load_ucsf_metadata(path: Path) -> dict[str, str]:
    """Load the UCSF source-ID to BraTS-ID metadata mapping."""

    mapping: dict[str, str] = {}
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            required = {"ID", "BraTS21 ID"}
            if not reader.fieldnames or not required.issubset(reader.fieldnames):
                raise OverlapAuditError(
                    "UCSF metadata is missing ID or BraTS21 ID columns"
                )
            for row in reader:
                raw_id = row.get("ID")
                raw_brats_id = row.get("BraTS21 ID")
                if raw_id is None or raw_brats_id is None or not raw_brats_id.strip():
                    continue
                source_id = normalize_ucsf_pdgm_id(raw_id)
                brats_id = normalize_identifier(raw_brats_id.strip(), Cohort.BRATS_2021)
                if source_id in mapping and mapping[source_id] != brats_id:
                    raise OverlapAuditError(
                        "UCSF metadata maps one source ID to multiple BraTS IDs"
                    )
                mapping[source_id] = brats_id
    except OSError as error:
        raise OverlapAuditError(f"cannot read UCSF metadata {path}") from error
    return mapping


def _crosscheck_ucsf_brats(
    metadata_map: Mapping[str, str],
    crosswalk_summary: Mapping[str, object],
    ucsf_ids: frozenset[str],
    crosswalk_pairs: Mapping[tuple[str, str], tuple[str, str]],
) -> tuple[dict[str, object], tuple[tuple[str, str], ...]]:
    """Compare metadata and crosswalk views without retaining IDs in summary."""

    metadata_brats = frozenset(metadata_map.values())
    primary = cast(frozenset[str], crosswalk_summary["ucsf_pdgm_primary_brats_ids"])
    additional = cast(
        frozenset[str], crosswalk_summary["ucsf_pdgm_additional_brats_ids"]
    )
    all_crosswalk = primary | additional
    metadata_source_ids = frozenset(metadata_map)
    metadata_in_dirs = metadata_source_ids & ucsf_ids
    pairs = tuple(
        sorted(
            (source_id, brats_id)
            for source_id, brats_id in metadata_map.items()
            if source_id in ucsf_ids and brats_id in all_crosswalk
        )
    )
    primary_pairs = tuple(
        sorted(pair for pair in crosswalk_pairs.values() if pair[0] in ucsf_ids)
    )
    summary: dict[str, object] = {
        "status": "disagreement" if metadata_brats != all_crosswalk else "agree",
        "metadata_nonblank_brats_id_count": len(metadata_brats),
        "metadata_source_id_count": len(metadata_source_ids),
        "metadata_source_ids_found_in_downloaded_directories": len(metadata_in_dirs),
        "crosswalk_ucsf_pdgm_brats_id_count": len(primary),
        "crosswalk_ucsf_pdgm_additional_brats_id_count": len(additional),
        "crosswalk_ucsf_all_brats_id_count": len(all_crosswalk),
        "metadata_ids_missing_from_crosswalk": len(metadata_brats - all_crosswalk),
        "crosswalk_ids_missing_from_metadata": len(all_crosswalk - metadata_brats),
        "primary_crosswalk_ids_missing_from_metadata": len(primary - metadata_brats),
        "additional_crosswalk_ids_missing_from_metadata": len(
            additional - metadata_brats
        ),
        "metadata_ids_missing_from_downloaded_directories": len(
            metadata_source_ids - ucsf_ids
        ),
        "crosswalk_primary_candidate_overlap_count": len(primary_pairs),
        "metadata_confirmed_overlap_count": len(pairs),
        "resolution_required": metadata_brats != all_crosswalk,
        "note": (
            "The matrix uses metadata-confirmed source IDs present in the "
            "downloaded UCSF directories for this cell; the alternate official "
            "crosswalk candidate count is retained above and must be resolved "
            "by T03.5 before external use."
        ),
    }
    return summary, pairs


def build_real_overlap_audit() -> tuple[dict[str, object], dict[str, object]]:
    """Build the aggregate and ID-level derived reports from local sources.

    This reads the real files through :mod:`glioma_selective.utils.paths` and
    never writes to ``data/raw``, ``data/external``, or ``data/private``.
    """

    utsw_ids = _normalize_set(
        Cohort.UTSW, (item.patient_id for item in utsw.discover_files())
    )
    ucsf_ids, ucsf_directory_finding = _normalize_ucsf_inventory_ids(
        item.patient_id for item in ucsf_pdgm.discover_files()
    )
    private_ids = paths.private_data_dir() / "external_cohort_ids"
    tcga_gbm_ids, gbm_finding = _load_tcga_ids(
        private_ids / "tcga_gbm_case_ids.json", Cohort.TCGA_GBM
    )
    tcga_lgg_ids, lgg_finding = _load_tcga_ids(
        private_ids / "tcga_lgg_case_ids.json", Cohort.TCGA_LGG
    )
    brats_ids, crosswalk_pairs, crosswalk_summary = _load_brats_crosswalk(
        private_ids / "brats2021_mapping_to_tcia.xlsx"
    )
    metadata_map = _load_ucsf_metadata(
        paths.external_data_dir() / "UCSF-PDGM-metadata_v2.csv"
    )
    crosscheck, ucsf_brats_pairs = _crosscheck_ucsf_brats(
        metadata_map, crosswalk_summary, ucsf_ids, crosswalk_pairs
    )

    source_notes = {
        "source_counts": {
            "utsw": len(utsw_ids),
            "ucsf_pdgm_raw_directories": ucsf_directory_finding["raw_directory_count"],
            "ucsf_pdgm_normalized_unique_patients": len(ucsf_ids),
            "brats2021_training_validation_crosswalk": len(brats_ids),
            "tcga_gbm": len(tcga_gbm_ids),
            "tcga_lgg": len(tcga_lgg_ids),
            "egd": 0,
        },
        "source_files": {
            "utsw": "paths.raw_data_dir()/utsw",
            "ucsf_pdgm": "paths.external_data_dir()/ucsf_pdgm",
            "ucsf_metadata": "paths.external_data_dir()/UCSF-PDGM-metadata_v2.csv",
            "brats2021": (
                "paths.private_data_dir()/external_cohort_ids/"
                "brats2021_mapping_to_tcia.xlsx"
            ),
            "tcga_gbm": (
                "paths.private_data_dir()/external_cohort_ids/tcga_gbm_case_ids.json"
            ),
            "tcga_lgg": (
                "paths.private_data_dir()/external_cohort_ids/tcga_lgg_case_ids.json"
            ),
        },
        "brats_test_set_caveat": BRATS_TEST_SET_CAVEAT,
        "egd_reason": EGD_NOT_AVAILABLE_REASON,
        "ucsf_pdgm_brats_crosscheck": crosscheck,
        "ucsf_directory_format_finding": ucsf_directory_finding,
        "brats_crosswalk_format_finding": {
            "bare_five_digit_brats_id_count": crosswalk_summary[
                "bare_five_digit_brats_id_count"
            ],
            "finding": crosswalk_summary["brats_format_finding"],
        },
        "tcga_id_format_findings": {
            "tcga_gbm": gbm_finding,
            "tcga_lgg": lgg_finding,
        },
    }
    aggregate = build_overlap_matrix(
        {
            Cohort.UTSW: utsw_ids,
            Cohort.UCSF_PDGM: ucsf_ids,
            Cohort.BRATS_2021: brats_ids,
            Cohort.TCGA_GBM: tcga_gbm_ids,
            Cohort.TCGA_LGG: tcga_lgg_ids,
        },
        pair_details={(Cohort.UCSF_PDGM, Cohort.BRATS_2021): ucsf_brats_pairs},
        source_notes=source_notes,
    )
    detail_pairs: dict[str, list[tuple[str, str]]] = {
        _pair_key(Cohort.UCSF_PDGM, Cohort.BRATS_2021): list(ucsf_brats_pairs),
    }
    for left in MATRIX_ROWS:
        for right in MATRIX_COLUMNS:
            if left == right or left == Cohort.EGD or right == Cohort.EGD:
                continue
            key = (left, right)
            if key == (Cohort.UCSF_PDGM, Cohort.BRATS_2021):
                continue
            if left in {
                Cohort.UTSW,
                Cohort.UCSF_PDGM,
            } and right in {
                Cohort.BRATS_2021,
                Cohort.TCGA_GBM,
                Cohort.TCGA_LGG,
                Cohort.UCSF_PDGM,
            }:
                left_ids = {
                    Cohort.UTSW: utsw_ids,
                    Cohort.UCSF_PDGM: ucsf_ids,
                }[left]
                right_ids = {
                    Cohort.BRATS_2021: brats_ids,
                    Cohort.TCGA_GBM: tcga_gbm_ids,
                    Cohort.TCGA_LGG: tcga_lgg_ids,
                    Cohort.UCSF_PDGM: ucsf_ids,
                }[right]
                detail_pairs[_pair_key(left, right)] = list(
                    _pair_details_for_intersection(left_ids, right_ids)
                )
    details = {
        "report_version": 1,
        "warning": (
            "This derived report contains raw-identifier-tier canonical IDs. "
            "Keep it under paths.manifests_dir(); never commit or publish it."
        ),
        "pairs": detail_pairs,
        "ucsf_pdgm_brats_crosscheck": crosscheck,
    }
    return aggregate, details


def write_overlap_report(
    report: Mapping[str, object], output_path: Path | None = None
) -> Path:
    """Write an aggregate-only matrix report under the manifests path."""

    destination = output_path or paths.manifests_dir() / "overlap_matrix.json"
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(
        json.dumps(report, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    return destination


def write_overlap_detail_report(
    report: Mapping[str, object], output_path: Path | None = None
) -> Path:
    """Write the derived ID-level report under the manifests path only."""

    destination = output_path or paths.manifests_dir() / "overlap_matrix_detail.json"
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(
        json.dumps(report, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    return destination
